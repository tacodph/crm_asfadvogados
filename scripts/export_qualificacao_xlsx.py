import json
from collections import Counter
from pathlib import Path

import openpyxl

wb = openpyxl.load_workbook(
    "documents/Qualificação de Atendimentos.xlsx",
    data_only=True,
)
print("sheets:", wb.sheetnames)
ws = wb["Agosto"] if "Agosto" in wb.sheetnames else wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print("headers:")
for i, h in enumerate(headers):
    print(i, repr(h))


def find(*parts):
    for i, h in enumerate(headers):
        if h and all(p.lower() in str(h).lower() for p in parts):
            return i
    return None


cols = {
    "data_entrada": find("data"),
    "nome": find("nome"),
    "ddd": find("ddd"),
    "whatsapp": find("whats"),
    "campanha": find("campanha"),
    "responsavel": find("respons"),
    "respondeu": find("respondeu"),
    "proposta": find("proposta"),
    "fu1": None,
    "fu2": None,
    "fu3": None,
    "status_atendimento": find("status", "atendimento"),
    "status_qualificacao": find("status", "qualifica"),
    "motivo_desqualificacao": find("motivo"),
    "continuidade": find("continuidade"),
    "observacoes": find("observa"),
}

# Follow-up columns if present
for i, h in enumerate(headers):
    if not h:
        continue
    hl = str(h).lower()
    if "fu1" in hl or "follow" in hl and "1" in hl:
        cols["fu1"] = i
    if "fu2" in hl or ("follow" in hl and "2" in hl):
        cols["fu2"] = i
    if "fu3" in hl or ("follow" in hl and "3" in hl):
        cols["fu3"] = i

print("col map", cols)

rows = []
att = Counter()
qual = Counter()
cont = Counter()

for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row):
        continue

    def val(k):
        i = cols[k]
        return None if i is None else row[i]

    item = {k: (None if val(k) in ("", None) else val(k)) for k in cols}
    for k in ["data_entrada", "fu1", "fu2", "fu3"]:
        v = item[k]
        if hasattr(v, "isoformat"):
            item[k] = v.date().isoformat() if hasattr(v, "date") else str(v)[:10]
    for k in ["ddd", "whatsapp"]:
        v = item[k]
        if v is not None:
            item[k] = str(v).split(".")[0]
    for k in [
        "respondeu",
        "proposta",
        "status_atendimento",
        "status_qualificacao",
        "motivo_desqualificacao",
        "continuidade",
        "observacoes",
        "nome",
        "campanha",
        "responsavel",
    ]:
        if item[k] is not None:
            item[k] = str(item[k]).strip()
    rows.append(item)
    if item["status_atendimento"]:
        att[item["status_atendimento"]] += 1
    if item["status_qualificacao"]:
        qual[item["status_qualificacao"]] += 1
    if item["continuidade"]:
        cont[item["continuidade"]] += 1

print("rows", len(rows))
print("status_atendimento", dict(att))
print("status_qualificacao", dict(qual))
print("continuidade", dict(cont))
Path("database/data/qualificacao-atendimentos-agosto.json").write_text(
    json.dumps(rows, ensure_ascii=False, indent=2),
    encoding="utf-8",
)
print("json updated")
